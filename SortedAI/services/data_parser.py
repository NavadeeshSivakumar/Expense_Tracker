import pdfplumber
import openpyxl
import io
from datetime import datetime

def extract_text_from_pdf(file_bytes: bytes) -> str:
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        text = ""
        for page in pdf.pages:
            text += page.extract_text() or ""
    return text

def extract_text_from_excel(file_bytes: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes))
    text = ""
    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        for row in ws.iter_rows(values_only=True):
            row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
            text += row_text + "\n"
    return text

def categorize(remarks: str) -> str:
    remarks_lower = remarks.lower()
    if any(k in remarks_lower for k in ["salary", "tata consultancy", "tcs", "neft-cit"]):
        return "Salary"
    if any(k in remarks_lower for k in ["swiggy", "zomato", "hungerbox", "popeyes", "kfc", "mcdonald", "food"]):
        return "Food"
    if any(k in remarks_lower for k in ["uber", "ola", "redbus", "bmrcl", "metro", "rapido"]):
        return "Transport"
    if any(k in remarks_lower for k in ["amazon", "flipkart", "myntra", "zepto", "blinkit"]):
        return "Shopping"
    if any(k in remarks_lower for k in ["netflix", "spotify", "bookmyshow", "bigtree", "hotstar"]):
        return "Entertainment"
    if any(k in remarks_lower for k in ["hospital", "pharmacy", "medical", "health", "clinic"]):
        return "Health"
    if any(k in remarks_lower for k in ["electricity", "water", "jio", "myjio", "airtel", "bsnl", "utility"]):
        return "Utilities"
    if any(k in remarks_lower for k in ["rent", "house", "pg", "hostel"]):
        return "Rent"
    if any(k in remarks_lower for k in ["freelance", "upwork", "fiverr"]):
        return "Freelance"
    return "Other"

def extract_transactions_from_excel(file_bytes: bytes) -> list:
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = workbook.active

    transactions = []
    header_found = False

    for row in ws.iter_rows(values_only=True):
        if not header_found:
            if row[1] == 'S No.':
                header_found = True
            continue

        if row[1] is None:
            continue
        if not isinstance(row[1], (int, float)):
            continue

        try:
            date_raw = row[3]
            remarks = str(row[5]) if row[5] else ""
            withdrawal = float(row[6]) if row[6] else 0
            deposit = float(row[7]) if row[7] else 0

            if withdrawal == 0 and deposit == 0:
                continue
            if date_raw is None:
                continue

            if isinstance(date_raw, str):
                date = datetime.strptime(date_raw, "%d/%m/%Y").strftime("%Y-%m-%d")
            else:
                date = date_raw.strftime("%Y-%m-%d")

            tx_type = "Income" if deposit > 0 else "Expense"
            amount = deposit if deposit > 0 else withdrawal
            note = remarks.split("/")[1] if "/" in remarks else remarks[:30]
            note = note.strip()
            category = categorize(remarks)

            transactions.append({
                "date": date,
                "type": tx_type,
                "category": category,
                "amount": amount,
                "note": note
            })

        except Exception as e:
            print(f"Skipped row {row[1]}: {e}")
            continue

    return transactions